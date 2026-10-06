"""Build default databases for THCS Phú Thịnh - Phân hiệu 1 (2026-2027):
- schools/truong-thcs-2-buoi.db
- data/sample_truong_thcs_2_buoi.db
- schools/truong-thcs.db
- data/sample_truong_thcs.db
Based on 'TKB_Sao_Luu_20261006_0959.xlsx', 'TKB_Tuan_6.xlsx', and 'Định lượng số tiết theo tuần năm học 2026_2027.xlsx'.
"""
from __future__ import annotations

import os
import shutil
import sys
import openpyxl

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from data import db, repository as repo
from io_excel.importer import import_xlsm
from io_excel.weekly_importer import import_weekly_curriculum_from_excel

SOURCE_EXCEL = "TKB_Sao_Luu_20261006_0959.xlsx"
WEEKLY_EXCEL = "Định lượng số tiết theo tuần năm học 2026_2027.xlsx"
TKB_TUAN_6_EXCEL = "TKB_Tuan_6.xlsx"
SCHOOL_NAME = "Trường THCS Phú Thịnh - Phân hiệu 1 (2026-2027)"

TARGET_PATHS = [
    os.path.join(os.path.dirname(__file__), "..", "schools", "truong-thcs-2-buoi.db"),
    os.path.join(os.path.dirname(__file__), "..", "data", "sample_truong_thcs_2_buoi.db"),
    os.path.join(os.path.dirname(__file__), "..", "schools", "truong-thcs.db"),
    os.path.join(os.path.dirname(__file__), "..", "data", "sample_truong_thcs.db"),
]


def load_tuan_6_cells(conn, excel_path: str) -> dict:
    classes = {c.name: c.class_id for c in repo.list_classes(conn)}
    subjects = {s.name: s.subject_id for s in repo.list_subjects(conn)}

    def norm_subj(name):
        name = name.strip()
        if name in subjects:
            return subjects[name]
        for s_name, sid in subjects.items():
            if s_name.lower() == name.lower():
                return sid
        return None

    wb = openpyxl.load_workbook(excel_path, data_only=True)
    ws = wb["TKB_Mon"]
    cells = {}
    for r in range(2, ws.max_row + 1):
        cname = ws.cell(r, 1).value
        sess = ws.cell(r, 2).value
        per = ws.cell(r, 3).value
        if not cname or cname not in classes:
            continue
        cid = classes[cname]
        for wd, col in enumerate(range(4, ws.max_column + 1), start=2):
            val = ws.cell(r, col).value
            if val:
                sid = norm_subj(str(val))
                if sid is not None:
                    cells[(cid, wd, sess, per)] = sid
    wb.close()
    return cells


def build_database(db_path: str):
    if os.path.exists(db_path):
        os.remove(db_path)

    os.makedirs(os.path.dirname(db_path), exist_ok=True)
    conn = db.get_connection(db_path)
    try:
        db.init_db(conn)
        rep_import = import_xlsm(conn, SOURCE_EXCEL)
        print(f"[{os.path.basename(db_path)}] import_xlsm: {rep_import.counts}")

        rep_weekly = import_weekly_curriculum_from_excel(conn, WEEKLY_EXCEL)
        print(f"[{os.path.basename(db_path)}] import_weekly: {rep_weekly['records_imported']} records across {rep_weekly['weeks_count']} weeks")

        is_2_buoi = "2_buoi" in db_path or "2-buoi" in db_path
        if is_2_buoi:
            repo.set_meta(conn, "school_name", "Trường THCS - Học 2 buổi (2026-2027)")
        else:
            repo.set_meta(conn, "school_name", SCHOOL_NAME)
        repo.set_meta(conn, "base_cap", "19")
        repo.set_meta(conn, "min_floor", "16")
        repo.set_meta(conn, "forbidden_off_cells", "2S")
        repo.set_meta(conn, "mandatory_morning_weekdays", "2")

        classes = conn.execute("SELECT class_id, name FROM classes").fetchall()
        k67 = [c for c in classes if c["name"].startswith(('6', '7'))]
        k89 = [c for c in classes if c["name"].startswith(('8', '9'))]

        if is_2_buoi:
            # Khung 2 buổi:
            # Khối 6-7 (29 tiết): T2-T4 (S1-S4, C1-C3 = 21t), T5 (S1-S4 = 4t), T6 (S1-S4 = 4t) -> 29t
            # Khối 8-9 (30 tiết): T2-T4 (S1-S4, C1-C3 = 21t), T5 (S1-S4 = 4t), T6 (S1-S5 = 5t) -> 30t
            for c in k67:
                cid = c["class_id"]
                conn.execute("DELETE FROM class_allowed_cells WHERE class_id = ?", (cid,))
                for w in (2, 3, 4):
                    for p in range(1, 5):
                        conn.execute("INSERT OR IGNORE INTO class_allowed_cells (class_id, weekday, session, period) VALUES (?, ?, 'S', ?)", (cid, str(w), p))
                    for p in range(1, 4):
                        conn.execute("INSERT OR IGNORE INTO class_allowed_cells (class_id, weekday, session, period) VALUES (?, ?, 'C', ?)", (cid, str(w), p))
                for p in range(1, 5):
                    conn.execute("INSERT OR IGNORE INTO class_allowed_cells (class_id, weekday, session, period) VALUES (?, '5', 'S', ?)", (cid, p))
                for p in range(1, 5):
                    conn.execute("INSERT OR IGNORE INTO class_allowed_cells (class_id, weekday, session, period) VALUES (?, '6', 'S', ?)", (cid, p))

                repo.set_frame_template(
                    conn, cid, morning_periods=4, afternoon_periods=3, study_sunday=False,
                    short_weekday=None, short_morning_periods=None, short_afternoon_periods=None
                )

            for c in k89:
                cid = c["class_id"]
                conn.execute("DELETE FROM class_allowed_cells WHERE class_id = ?", (cid,))
                for w in (2, 3, 4):
                    for p in range(1, 5):
                        conn.execute("INSERT OR IGNORE INTO class_allowed_cells (class_id, weekday, session, period) VALUES (?, ?, 'S', ?)", (cid, str(w), p))
                    for p in range(1, 4):
                        conn.execute("INSERT OR IGNORE INTO class_allowed_cells (class_id, weekday, session, period) VALUES (?, ?, 'C', ?)", (cid, str(w), p))
                for p in range(1, 5):
                    conn.execute("INSERT OR IGNORE INTO class_allowed_cells (class_id, weekday, session, period) VALUES (?, '5', 'S', ?)", (cid, p))
                for p in range(1, 6):
                    conn.execute("INSERT OR IGNORE INTO class_allowed_cells (class_id, weekday, session, period) VALUES (?, '6', 'S', ?)", (cid, p))

                repo.set_frame_template(
                    conn, cid, morning_periods=4, afternoon_periods=3, study_sunday=False,
                    short_weekday=6, short_morning_periods=5, short_afternoon_periods=None
                )
        else:
            # Khung 1 buổi: K67 = 29 periods, K89 = 30 periods

            for c in k67:
                cid = c["class_id"]
                conn.execute("DELETE FROM class_allowed_cells WHERE class_id = ? AND weekday = '7' AND session = 'S' AND period = 5", (cid,))
                for w in range(2, 7):
                    for p in range(1, 6):
                        conn.execute("INSERT OR IGNORE INTO class_allowed_cells (class_id, weekday, session, period) VALUES (?, ?, 'S', ?)", (cid, str(w), p))
                for p in range(1, 5):
                    conn.execute("INSERT OR IGNORE INTO class_allowed_cells (class_id, weekday, session, period) VALUES (?, '7', 'S', ?)", (cid, p))

            for c in k89:
                cid = c["class_id"]
                for w in range(2, 8):
                    for p in range(1, 6):
                        conn.execute("INSERT OR IGNORE INTO class_allowed_cells (class_id, weekday, session, period) VALUES (?, ?, 'S', ?)", (cid, str(w), p))

        # Import Week 6 official timetable if file exists
        if os.path.exists(TKB_TUAN_6_EXCEL):
            cells_w6 = load_tuan_6_cells(conn, TKB_TUAN_6_EXCEL)
            if cells_w6:
                run_id = repo.save_run(
                    conn,
                    week_no=6,
                    seed=2026,
                    parity="C",
                    cells_changed=0,
                    cells_total=len(cells_w6),
                    succeeded=True,
                    message="Thời khóa biểu chính thức Tuần 6 (Đã nhập từ file TKB_Tuan_6.xlsx)",
                )
                repo.save_tkb_result(conn, run_id, cells_w6)
                repo.bulk_replace_tkb_nhap(conn, cells_w6)
                print(f"[{os.path.basename(db_path)}] Saved Week 6 official timetable: {len(cells_w6)} cells (run_id={run_id})")

        conn.commit()

        # Verification
        classes_all = repo.list_classes(conn)
        subjects = repo.list_subjects(conn)
        teachers = repo.list_teachers(conn)
        saved_weeks = repo.list_saved_weeks(conn)
        print(f"[{os.path.basename(db_path)}] Verification: {len(classes_all)} classes, {len(subjects)} subjects, {len(teachers)} teachers, saved_weeks={saved_weeks}")
    finally:
        conn.close()


def main():
    for p in TARGET_PATHS:
        build_database(p)
    print("All databases created and configured successfully!")


if __name__ == "__main__":
    main()
