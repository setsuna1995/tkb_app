"""Script to create and populate the school database and standard Excel file
for: TRƯỜNG THCS ĐỖ CẬN - PHÂN HIỆU (Năm học 2026-2027)
from the official teaching assignment schedule image.
"""
from __future__ import annotations

import os
import sys

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

import sqlite3
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from core.models import ROLE_GDTC, ROLE_HDTN, ROLE_NANG, ROLE_NANG_KEP, ROLE_THUONG
from data import db
from data import repository as repo
from io_excel.weekly_importer import import_weekly_curriculum_from_excel

DB_PATH = os.path.join("schools", "thcs-do-can-phan-hieu.db")
EXCEL_PATH = "THCS_Do_Can_Phan_Hieu.xlsx"

# 1. Classes: 20 classes across 4 grades
CLASSES_DATA = [
    # Grade 6
    ("6A7", 0), ("6A8", 1), ("6A9", 2), ("6A10", 3), ("6A11", 4),
    # Grade 7
    ("7A7", 5), ("7A8", 6), ("7A9", 7), ("7A10", 8), ("7A11", 9),
    # Grade 8
    ("8A7", 10), ("8A8", 11), ("8A9", 12), ("8A10", 13), ("8A11", 14),
    # Grade 9
    ("9A8", 15), ("9A9", 16), ("9A10", 17), ("9A11", 18), ("9A12", 19),
]

# 2. Subjects: Standard 16 subjects GDPT 2018
SUBJECTS_DATA = [
    ("Toán học", ROLE_NANG_KEP, 0),
    ("Ngữ văn", ROLE_NANG_KEP, 1),
    ("Tiếng Anh", ROLE_NANG, 2),
    ("Khoa học tự nhiên (Vật lý)", ROLE_THUONG, 3),
    ("Khoa học tự nhiên (Hóa học)", ROLE_THUONG, 4),
    ("Khoa học tự nhiên (Sinh học)", ROLE_THUONG, 5),
    ("Lịch sử và Địa Lý (Lịch sử)", ROLE_THUONG, 6),
    ("Lịch sử và Địa Lý (Địa lý)", ROLE_THUONG, 7),
    ("GDCD", ROLE_THUONG, 8),
    ("Công nghệ", ROLE_THUONG, 9),
    ("Tin học", ROLE_THUONG, 10),
    ("Giáo dục thể chất", ROLE_GDTC, 11),
    ("Nội dung giáo dục của địa phương", ROLE_THUONG, 12),
    ("Hoạt động trải nghiệm, hướng nghiệp", ROLE_HDTN, 13),
    ("Nghệ thuật (Âm nhạc)", ROLE_THUONG, 14),
    ("Nghệ thuật (Mỹ thuật)", ROLE_THUONG, 15),
]

# 3. Teachers: 41 records
# (stt, name, subject_specialty, gvcn_class, role, role_reduction, thuc_day, tong, thua_thieu)
TEACHERS_RAW = [
    (1, "Đỗ Thị Sa", "Tiếng Anh", None, "Hiệu trưởng", 17, 2, 2, None),
    (2, "Đoàn Tiến Quân", "Sinh-Kỹ", None, "Phó hiệu trưởng", 15, 2, 2, None),
    (3, "Phan Thị Thùy Linh", "Văn-Sử", "8A7", "Tổ trưởng", 3, 12, 19, 0),
    (4, "Nguyễn Thị Thu Hương", "Văn-Sử", "8A9", "GVCN", 4, 17, 21, 2),
    (5, "Đồng Thị Hiền", "Văn-Sử", "7A8", "GVCN", 4, 16, 20, 1),
    (6, "Hoàng Thị Phương Duyên", "Văn-Sử", "6A8", "GVCN", 4, 17, 21, 2),
    (7, "Ngô Thị Xuân Hương", "Văn", "8A8", "GVCN", 4, 17, 21, 2),
    (8, "Bùi Thị Mỹ Hạnh", "Văn-Sử", "7A9", "GVCN", 4, 17, 21, 2),
    (9, "Triệu Thị Luyến", "Văn", "9A12", "GVCN", 4, 16, 20, 1),
    (10, "Hoàng Thị Kiều Loan", "Văn", None, "HĐ235", 0, 19, 19, 0),
    (11, "Bàng Thị Phượng", "Văn", None, "HĐ235", 0, 19, 19, 0),
    (12, "Nông Thị Trà", "Văn", None, "HĐ235", 0, 19, 19, 0),
    (13, "Bùi Thị Hòa", "Lịch Sử", "9A9", "GVCN", 4, 17, 21, 2),
    (14, "Lại Thị Thanh Huyền", "Toán", "6A7", "Tổ phó", 1, 15, 20, 1),
    (15, "Nguyễn Thị Hà", "Toán", "7A7", "GVCN", 4, 16, 20, 1),
    (16, "Mẫn Xuân Thắng", "Toán-Lí", None, "", 0, 19, 19, 0),
    (17, "Phùng Thị Lan", "Toán", "9A10", "GVCN", 4, 15, 19, 0),
    (18, "Hoàng Thị Thanh Thắm", "GDTC", "7A10", "GVCN", 4, 15, 19, 0),
    (19, "Lê Thị Thu Hằng", "Toán", "6A11", "HĐ235 / GVCN", 4, 16, 20, 1),
    (20, "Nguyễn Thị Hiền", "Toán-Tin", None, "HĐ235", 0, 19, 19, 0),
    (21, "Lương Thị Thu Hiền", "Tin học", None, "HĐ235", 0, 19, 19, 0),
    (22, "Hà Thị Hải Yến", "Sinh-Kỹ", "8A10", "Tổ phó", 1, 15, 20, 1),
    (23, "Trần Thị Liễu", "Sinh-Hoá", None, "HĐK", 0, 20, 20, 1),
    (24, "Trần Thị Quỳnh Hằng", "Sinh", None, "HĐK", 0, 8, 8, 0),
    (25, "Nguyễn Ánh Hồng", "Sinh", None, "HĐK", 0, 8, 8, 0),
    (26, "Vũ Thị Bích Hằng", "Sinh(Hoá-Địa)", None, "", 0, 19, 19, 0),
    (27, "Chu Diệu Linh", "Địa", "9A11", "GVCN", 4, 15, 19, 0),
    (28, "Trương Thị Thu Hương", "Lí-CN", "7A11", "GVCN", 4, 15, 19, 0),
    (29, "Đỗ Thị Thu Hiền", "Lí", None, "", 0, 19, 19, 0),
    (30, "Lê Văn Hậu", "GDTC", None, "", 0, 20, 20, 1),
    (31, "GV môn GDTC", "GDTC", None, "", 0, 18, 18, -1),
    (32, "Bùi Đức Anh", "GDTC", None, "Trường chính", 0, 2, 2, 0),
    (33, "Lý Thị Duyên", "MT-Đội", "6A10", "GVCN", 4, 15, 19, 0),
    (34, "Lê Thị Minh Anh", "Mĩ thuật", None, "HĐK", 0, 4, 4, 0),
    (35, "Ngô Thuý Vân", "MT-Đội", None, "Tổng phụ trách", 8, 1, 1, 0),
    (36, "Nguyễn Thị Luyến", "Âm Nhạc", None, "", 0, 20, 20, 1),
    (37, "Nguyễn Đình Tiến", "Tiếng Anh", "9A8", "GVCN", 4, 15, 19, 0),
    (38, "Phan Thị Minh", "Tiếng Anh", "6A9", "Tổ phó", 1, 15, 20, 1),
    (39, "Trần Đức Ngân", "Tiếng Anh", "8A11", "GVCN", 4, 17, 21, 2),
    (40, "Nguyễn Thị Hồng Minh", "Tiếng Anh", None, "HĐ235", 0, 19, 19, 0),
    (41, "Bàng Như Quỳnh", "Tiếng Anh", None, "Trường chính", 0, 1, 1, 0),
]

# 4. Detailed Assignments (Teacher -> list of (subject, class, periods_per_week, note))
ASSIGNMENTS_FLAT = [
    # 1. Đỗ Thị Sa (HT)
    (1, "Đỗ Thị Sa", "Tiếng Anh", "6A7", 1, "Dạy chung lớp 6A7"),
    (1, "Đỗ Thị Sa", "Tiếng Anh", "6A8", 1, "Dạy chung lớp 6A8"),
    # 2. Đoàn Tiến Quân (PHT)
    (2, "Đoàn Tiến Quân", "Khoa học tự nhiên (Sinh học)", "7A7", 2, "Dạy PHT"),
    # 3. Phan Thị Thùy Linh (TTCM, CN 8A7)
    (3, "Phan Thị Thùy Linh", "Ngữ văn", "8A7", 4, ""),
    (3, "Phan Thị Thùy Linh", "Ngữ văn", "9A9", 4, ""),
    (3, "Phan Thị Thùy Linh", "Nội dung giáo dục của địa phương", "6A7", 2, ""),
    (3, "Phan Thị Thùy Linh", "BDHSG Văn", "Khối 8/9", 2, "BDHSG (ngoài TKB chính)"),
    # 4. Nguyễn Thị Thu Hương (CN 8A9)
    (4, "Nguyễn Thị Thu Hương", "Ngữ văn", "8A9", 4, ""),
    (4, "Nguyễn Thị Thu Hương", "Ngữ văn", "9A11", 4, ""),
    (4, "Nguyễn Thị Thu Hương", "Lịch sử và Địa Lý (Lịch sử)", "7A9", 2, ""),
    (4, "Nguyễn Thị Thu Hương", "Lịch sử và Địa Lý (Lịch sử)", "7A10", 2, ""),
    (4, "Nguyễn Thị Thu Hương", "Lịch sử và Địa Lý (Lịch sử)", "6A7", 1, ""),
    (4, "Nguyễn Thị Thu Hương", "Lịch sử và Địa Lý (Lịch sử)", "6A8", 1, ""),
    (4, "Nguyễn Thị Thu Hương", "Lịch sử và Địa Lý (Lịch sử)", "6A9", 1, ""),
    (4, "Nguyễn Thị Thu Hương", "Lịch sử và Địa Lý (Lịch sử)", "6A10", 1, ""),
    (4, "Nguyễn Thị Thu Hương", "Lịch sử và Địa Lý (Lịch sử)", "6A11", 1, ""),
    # 5. Đồng Thị Hiền (CN 7A8)
    (5, "Đồng Thị Hiền", "Ngữ văn", "7A8", 4, ""),
    (5, "Đồng Thị Hiền", "Ngữ văn", "9A8", 4, ""),
    (5, "Đồng Thị Hiền", "Lịch sử và Địa Lý (Lịch sử)", "9A8", 2, ""),
    (5, "Đồng Thị Hiền", "Lịch sử và Địa Lý (Lịch sử)", "9A10", 2, ""),
    (5, "Đồng Thị Hiền", "Nội dung giáo dục của địa phương", "9A11", 1, ""),
    (5, "Đồng Thị Hiền", "Nội dung giáo dục của địa phương", "9A12", 1, ""),
    (5, "Đồng Thị Hiền", "BDHSG Sử", "Khối 9", 2, "BDHSG (ngoài TKB chính)"),
    # 6. Hoàng Thị Phương Duyên (CN 6A8)
    (6, "Hoàng Thị Phương Duyên", "Ngữ văn", "6A8", 4, ""),
    (6, "Hoàng Thị Phương Duyên", "Ngữ văn", "6A9", 4, ""),
    (6, "Hoàng Thị Phương Duyên", "Ngữ văn", "8A10", 4, ""),
    (6, "Hoàng Thị Phương Duyên", "GDCD", "9A8", 1, ""),
    (6, "Hoàng Thị Phương Duyên", "GDCD", "9A9", 1, ""),
    (6, "Hoàng Thị Phương Duyên", "GDCD", "9A10", 1, ""),
    (6, "Hoàng Thị Phương Duyên", "GDCD", "9A11", 1, ""),
    (6, "Hoàng Thị Phương Duyên", "GDCD", "9A12", 1, ""),
    # 7. Ngô Thị Xuân Hương (CN 8A8)
    (7, "Ngô Thị Xuân Hương", "Ngữ văn", "8A8", 4, ""),
    (7, "Ngô Thị Xuân Hương", "Ngữ văn", "8A11", 4, ""),
    (7, "Ngô Thị Xuân Hương", "Ngữ văn", "6A7", 4, ""),
    (7, "Ngô Thị Xuân Hương", "Nội dung giáo dục của địa phương", "8A7", 1, ""),
    (7, "Ngô Thị Xuân Hương", "Nội dung giáo dục của địa phương", "8A8", 1, ""),
    (7, "Ngô Thị Xuân Hương", "Nội dung giáo dục của địa phương", "8A9", 1, ""),
    (7, "Ngô Thị Xuân Hương", "Nội dung giáo dục của địa phương", "8A10", 1, ""),
    (7, "Ngô Thị Xuân Hương", "Nội dung giáo dục của địa phương", "8A11", 1, ""),
    # 8. Bùi Thị Mỹ Hạnh (CN 7A9)
    (8, "Bùi Thị Mỹ Hạnh", "Ngữ văn", "7A9", 4, ""),
    (8, "Bùi Thị Mỹ Hạnh", "Ngữ văn", "7A10", 4, ""),
    (8, "Bùi Thị Mỹ Hạnh", "Ngữ văn", "9A10", 4, ""),
    (8, "Bùi Thị Mỹ Hạnh", "GDCD", "8A7", 1, ""),
    (8, "Bùi Thị Mỹ Hạnh", "GDCD", "8A8", 1, ""),
    (8, "Bùi Thị Mỹ Hạnh", "GDCD", "8A9", 1, ""),
    (8, "Bùi Thị Mỹ Hạnh", "GDCD", "8A10", 1, ""),
    (8, "Bùi Thị Mỹ Hạnh", "GDCD", "8A11", 1, ""),
    # 9. Triệu Thị Luyến (CN 9A12)
    (9, "Triệu Thị Luyến", "Ngữ văn", "9A12", 4, ""),
    (9, "Triệu Thị Luyến", "Ngữ văn", "7A7", 4, ""),
    (9, "Triệu Thị Luyến", "Nội dung giáo dục của địa phương", "7A7", 1, ""),
    (9, "Triệu Thị Luyến", "Nội dung giáo dục của địa phương", "7A8", 1, ""),
    (9, "Triệu Thị Luyến", "Nội dung giáo dục của địa phương", "7A9", 1, ""),
    (9, "Triệu Thị Luyến", "Nội dung giáo dục của địa phương", "7A10", 1, ""),
    (9, "Triệu Thị Luyến", "Nội dung giáo dục của địa phương", "7A11", 1, ""),
    (9, "Triệu Thị Luyến", "Nội dung giáo dục của địa phương", "9A8", 1, ""),
    (9, "Triệu Thị Luyến", "Nội dung giáo dục của địa phương", "9A9", 1, ""),
    (9, "Triệu Thị Luyến", "Nội dung giáo dục của địa phương", "9A10", 1, ""),
    # 10. Hoàng Thị Kiều Loan (HĐ235)
    (10, "Hoàng Thị Kiều Loan", "Ngữ văn", "7A11", 4, ""),
    (10, "Hoàng Thị Kiều Loan", "GDCD", "6A7", 1, ""),
    (10, "Hoàng Thị Kiều Loan", "GDCD", "6A8", 1, ""),
    (10, "Hoàng Thị Kiều Loan", "GDCD", "6A9", 1, ""),
    (10, "Hoàng Thị Kiều Loan", "GDCD", "6A10", 1, ""),
    (10, "Hoàng Thị Kiều Loan", "GDCD", "6A11", 1, ""),
    (10, "Hoàng Thị Kiều Loan", "Hoạt động trải nghiệm, hướng nghiệp", "8A7", 3, ""),
    (10, "Hoàng Thị Kiều Loan", "Hoạt động trải nghiệm, hướng nghiệp", "8A8", 3, ""),
    (10, "Hoàng Thị Kiều Loan", "Hoạt động trải nghiệm, hướng nghiệp", "8A9", 3, ""),
    (10, "Hoàng Thị Kiều Loan", "Công nghệ", "7A7", 1, ""),
    # 11. Bàng Thị Phượng (HĐ235)
    (11, "Bàng Thị Phượng", "Ngữ văn", "6A11", 4, ""),
    (11, "Bàng Thị Phượng", "GDCD", "7A7", 1, ""),
    (11, "Bàng Thị Phượng", "GDCD", "7A8", 1, ""),
    (11, "Bàng Thị Phượng", "GDCD", "7A9", 1, ""),
    (11, "Bàng Thị Phượng", "GDCD", "7A10", 1, ""),
    (11, "Bàng Thị Phượng", "GDCD", "7A11", 1, ""),
    (11, "Bàng Thị Phượng", "Hoạt động trải nghiệm, hướng nghiệp", "6A10", 3, ""),
    (11, "Bàng Thị Phượng", "Hoạt động trải nghiệm, hướng nghiệp", "6A11", 3, ""),
    (11, "Bàng Thị Phượng", "Nội dung giáo dục của địa phương", "6A8", 1, ""),
    (11, "Bàng Thị Phượng", "Nội dung giáo dục của địa phương", "6A9", 1, ""),
    (11, "Bàng Thị Phượng", "Nội dung giáo dục của địa phương", "6A10", 1, ""),
    (11, "Bàng Thị Phượng", "Nội dung giáo dục của địa phương", "6A11", 1, ""),
    # 12. Nông Thị Trà (HĐ235)
    (12, "Nông Thị Trà", "Ngữ văn", "6A10", 4, ""),
    (12, "Nông Thị Trà", "Hoạt động trải nghiệm, hướng nghiệp", "9A8", 3, ""),
    (12, "Nông Thị Trà", "Hoạt động trải nghiệm, hướng nghiệp", "9A9", 3, ""),
    (12, "Nông Thị Trà", "Hoạt động trải nghiệm, hướng nghiệp", "9A11", 3, ""),
    (12, "Nông Thị Trà", "Hoạt động trải nghiệm, hướng nghiệp", "9A12", 3, ""),
    (12, "Nông Thị Trà", "Hoạt động trải nghiệm, hướng nghiệp", "7A11", 3, ""),
    # 13. Bùi Thị Hòa (CN 9A9)
    (13, "Bùi Thị Hòa", "Lịch sử và Địa Lý (Lịch sử)", "9A9", 2, ""),
    (13, "Bùi Thị Hòa", "Lịch sử và Địa Lý (Lịch sử)", "9A11", 2, ""),
    (13, "Bùi Thị Hòa", "Lịch sử và Địa Lý (Lịch sử)", "9A12", 2, ""),
    (13, "Bùi Thị Hòa", "Lịch sử và Địa Lý (Lịch sử)", "8A7", 1, ""),
    (13, "Bùi Thị Hòa", "Lịch sử và Địa Lý (Lịch sử)", "8A8", 1, ""),
    (13, "Bùi Thị Hòa", "Lịch sử và Địa Lý (Lịch sử)", "8A9", 1, ""),
    (13, "Bùi Thị Hòa", "Lịch sử và Địa Lý (Lịch sử)", "8A10", 1, ""),
    (13, "Bùi Thị Hòa", "Lịch sử và Địa Lý (Lịch sử)", "8A11", 1, ""),
    (13, "Bùi Thị Hòa", "Lịch sử và Địa Lý (Lịch sử)", "7A7", 2, ""),
    (13, "Bùi Thị Hòa", "Lịch sử và Địa Lý (Lịch sử)", "7A8", 2, ""),
    (13, "Bùi Thị Hòa", "Lịch sử và Địa Lý (Lịch sử)", "7A11", 2, ""),
    # 14. Lại Thị Thanh Huyền (CN 6A7, TPCM)
    (14, "Lại Thị Thanh Huyền", "Toán học", "6A7", 4, ""),
    (14, "Lại Thị Thanh Huyền", "Toán học", "6A8", 4, ""),
    (14, "Lại Thị Thanh Huyền", "Toán học", "9A9", 4, ""),
    (14, "Lại Thị Thanh Huyền", "Hoạt động trải nghiệm, hướng nghiệp", "6A7", 3, ""),
    # 15. Nguyễn Thị Hà (CN 7A7)
    (15, "Nguyễn Thị Hà", "Toán học", "7A7", 4, ""),
    (15, "Nguyễn Thị Hà", "Toán học", "7A8", 4, ""),
    (15, "Nguyễn Thị Hà", "Toán học", "9A8", 4, ""),
    (15, "Nguyễn Thị Hà", "Toán học", "9A12", 4, ""),
    # 16. Mẫn Xuân Thắng
    (16, "Mẫn Xuân Thắng", "Toán học", "8A9", 4, ""),
    (16, "Mẫn Xuân Thắng", "Công nghệ", "9A8", 2, ""),
    (16, "Mẫn Xuân Thắng", "Công nghệ", "9A9", 2, ""),
    (16, "Mẫn Xuân Thắng", "Công nghệ", "9A10", 2, ""),
    (16, "Mẫn Xuân Thắng", "Công nghệ", "9A11", 2, ""),
    (16, "Mẫn Xuân Thắng", "Công nghệ", "9A12", 2, ""),
    (16, "Mẫn Xuân Thắng", "Công nghệ", "8A7", 1, ""),
    (16, "Mẫn Xuân Thắng", "Công nghệ", "8A8", 1, ""),
    (16, "Mẫn Xuân Thắng", "Công nghệ", "8A9", 1, ""),
    (16, "Mẫn Xuân Thắng", "Công nghệ", "8A10", 1, ""),
    (16, "Mẫn Xuân Thắng", "Công nghệ", "8A11", 1, ""),
    # 17. Phùng Thị Lan (CN 9A10)
    (17, "Phùng Thị Lan", "Toán học", "9A10", 4, ""),
    (17, "Phùng Thị Lan", "Toán học", "7A9", 4, ""),
    (17, "Phùng Thị Lan", "Toán học", "7A11", 4, ""),
    (17, "Phùng Thị Lan", "Hoạt động trải nghiệm, hướng nghiệp", "9A10", 3, ""),
    # 18. Hoàng Thị Thanh Thắm (CN 7A10)
    (18, "Hoàng Thị Thanh Thắm", "Toán học", "7A10", 4, ""),
    (18, "Hoàng Thị Thanh Thắm", "Toán học", "8A7", 4, ""),
    (18, "Hoàng Thị Thanh Thắm", "Toán học", "8A8", 4, ""),
    (18, "Hoàng Thị Thanh Thắm", "Hoạt động trải nghiệm, hướng nghiệp", "7A10", 3, ""),
    # 19. Lê Thị Thu Hằng (HĐ235, CN 6A11)
    (19, "Lê Thị Thu Hằng", "Toán học", "6A9", 4, ""),
    (19, "Lê Thị Thu Hằng", "Toán học", "6A10", 4, ""),
    (19, "Lê Thị Thu Hằng", "Toán học", "6A11", 4, ""),
    (19, "Lê Thị Thu Hằng", "Toán học", "9A11", 4, ""),
    # 20. Nguyễn Thị Hiền (HĐ235)
    (20, "Nguyễn Thị Hiền", "Toán học", "8A10", 4, ""),
    (20, "Nguyễn Thị Hiền", "Toán học", "8A11", 4, ""),
    (20, "Nguyễn Thị Hiền", "Tin học", "9A8", 1, ""),
    (20, "Nguyễn Thị Hiền", "Tin học", "9A9", 1, ""),
    (20, "Nguyễn Thị Hiền", "Tin học", "9A10", 1, ""),
    (20, "Nguyễn Thị Hiền", "Tin học", "9A11", 1, ""),
    (20, "Nguyễn Thị Hiền", "Tin học", "9A12", 1, ""),
    (20, "Nguyễn Thị Hiền", "Hoạt động trải nghiệm, hướng nghiệp", "7A7", 3, ""),
    (20, "Nguyễn Thị Hiền", "Hoạt động trải nghiệm, hướng nghiệp", "7A8", 3, ""),
    # 21. Lương Thị Thu Hiền (HĐ235)
    (21, "Lương Thị Thu Hiền", "Tin học", "6A7", 1, ""),
    (21, "Lương Thị Thu Hiền", "Tin học", "6A8", 1, ""),
    (21, "Lương Thị Thu Hiền", "Tin học", "6A9", 1, ""),
    (21, "Lương Thị Thu Hiền", "Tin học", "6A10", 1, ""),
    (21, "Lương Thị Thu Hiền", "Tin học", "6A11", 1, ""),
    (21, "Lương Thị Thu Hiền", "Tin học", "7A7", 1, ""),
    (21, "Lương Thị Thu Hiền", "Tin học", "7A8", 1, ""),
    (21, "Lương Thị Thu Hiền", "Tin học", "7A9", 1, ""),
    (21, "Lương Thị Thu Hiền", "Tin học", "7A10", 1, ""),
    (21, "Lương Thị Thu Hiền", "Tin học", "7A11", 1, ""),
    (21, "Lương Thị Thu Hiền", "Tin học", "8A7", 1, ""),
    (21, "Lương Thị Thu Hiền", "Tin học", "8A8", 1, ""),
    (21, "Lương Thị Thu Hiền", "Tin học", "8A9", 1, ""),
    (21, "Lương Thị Thu Hiền", "Tin học", "8A10", 1, ""),
    (21, "Lương Thị Thu Hiền", "Tin học", "8A11", 1, ""),
    (21, "Lương Thị Thu Hiền", "Công nghệ", "7A8", 1, ""),
    (21, "Lương Thị Thu Hiền", "Công nghệ", "7A9", 1, ""),
    (21, "Lương Thị Thu Hiền", "Công nghệ", "7A10", 1, ""),
    (21, "Lương Thị Thu Hiền", "Công nghệ", "7A11", 1, ""),
    # 22. Hà Thị Hải Yến (CN 8A10, TPCM)
    (22, "Hà Thị Hải Yến", "Khoa học tự nhiên (Sinh học)", "9A8", 1, ""),
    (22, "Hà Thị Hải Yến", "Khoa học tự nhiên (Sinh học)", "9A9", 1, ""),
    (22, "Hà Thị Hải Yến", "Khoa học tự nhiên (Sinh học)", "9A10", 1, ""),
    (22, "Hà Thị Hải Yến", "Khoa học tự nhiên (Sinh học)", "9A11", 1, ""),
    (22, "Hà Thị Hải Yến", "Khoa học tự nhiên (Sinh học)", "9A12", 1, ""),
    (22, "Hà Thị Hải Yến", "Khoa học tự nhiên (Sinh học)", "8A10", 2, ""),
    (22, "Hà Thị Hải Yến", "Khoa học tự nhiên (Sinh học)", "6A7", 1, ""),
    (22, "Hà Thị Hải Yến", "Khoa học tự nhiên (Sinh học)", "6A8", 1, ""),
    (22, "Hà Thị Hải Yến", "Khoa học tự nhiên (Sinh học)", "6A9", 1, ""),
    (22, "Hà Thị Hải Yến", "Khoa học tự nhiên (Sinh học)", "6A10", 1, ""),
    (22, "Hà Thị Hải Yến", "Khoa học tự nhiên (Sinh học)", "6A11", 1, ""),
    (22, "Hà Thị Hải Yến", "Hoạt động trải nghiệm, hướng nghiệp", "8A10", 3, ""),
    # 23. Trần Thị Liễu (HĐK)
    (23, "Trần Thị Liễu", "Khoa học tự nhiên (Hóa học)", "7A7", 1, ""),
    (23, "Trần Thị Liễu", "Khoa học tự nhiên (Hóa học)", "7A8", 1, ""),
    (23, "Trần Thị Liễu", "Khoa học tự nhiên (Hóa học)", "7A9", 1, ""),
    (23, "Trần Thị Liễu", "Khoa học tự nhiên (Hóa học)", "7A10", 1, ""),
    (23, "Trần Thị Liễu", "Khoa học tự nhiên (Hóa học)", "7A11", 1, ""),
    (23, "Trần Thị Liễu", "Khoa học tự nhiên (Hóa học)", "8A7", 1, ""),
    (23, "Trần Thị Liễu", "Khoa học tự nhiên (Hóa học)", "8A8", 1, ""),
    (23, "Trần Thị Liễu", "Khoa học tự nhiên (Hóa học)", "8A9", 1, ""),
    (23, "Trần Thị Liễu", "Khoa học tự nhiên (Hóa học)", "8A10", 1, ""),
    (23, "Trần Thị Liễu", "Khoa học tự nhiên (Hóa học)", "8A11", 1, ""),
    (23, "Trần Thị Liễu", "Khoa học tự nhiên (Hóa học)", "9A8", 1, ""),
    (23, "Trần Thị Liễu", "Khoa học tự nhiên (Hóa học)", "9A9", 1, ""),
    (23, "Trần Thị Liễu", "Khoa học tự nhiên (Hóa học)", "9A10", 1, ""),
    (23, "Trần Thị Liễu", "Khoa học tự nhiên (Hóa học)", "9A11", 1, ""),
    (23, "Trần Thị Liễu", "Khoa học tự nhiên (Hóa học)", "9A12", 1, ""),
    (23, "Trần Thị Liễu", "Khoa học tự nhiên (Hóa học)", "6A7", 1, ""),
    (23, "Trần Thị Liễu", "Khoa học tự nhiên (Hóa học)", "6A8", 1, ""),
    (23, "Trần Thị Liễu", "Khoa học tự nhiên (Hóa học)", "6A9", 1, ""),
    (23, "Trần Thị Liễu", "Khoa học tự nhiên (Hóa học)", "6A10", 1, ""),
    (23, "Trần Thị Liễu", "Khoa học tự nhiên (Hóa học)", "6A11", 1, ""),
    # 24. Trần Thị Quỳnh Hằng (HĐK)
    (24, "Trần Thị Quỳnh Hằng", "Khoa học tự nhiên (Sinh học)", "8A7", 2, ""),
    (24, "Trần Thị Quỳnh Hằng", "Khoa học tự nhiên (Sinh học)", "8A8", 2, ""),
    (24, "Trần Thị Quỳnh Hằng", "Khoa học tự nhiên (Sinh học)", "8A9", 2, ""),
    (24, "Trần Thị Quỳnh Hằng", "Khoa học tự nhiên (Sinh học)", "8A11", 2, ""),
    # 25. Nguyễn Ánh Hồng (HĐK)
    (25, "Nguyễn Ánh Hồng", "Khoa học tự nhiên (Sinh học)", "7A8", 2, ""),
    (25, "Nguyễn Ánh Hồng", "Khoa học tự nhiên (Sinh học)", "7A9", 2, ""),
    (25, "Nguyễn Ánh Hồng", "Khoa học tự nhiên (Sinh học)", "7A10", 2, ""),
    (25, "Nguyễn Ánh Hồng", "Khoa học tự nhiên (Sinh học)", "7A11", 2, ""),
    # 26. Vũ Thị Bích Hằng
    (26, "Vũ Thị Bích Hằng", "Lịch sử và Địa Lý (Địa lý)", "6A7", 2, ""),
    (26, "Vũ Thị Bích Hằng", "Lịch sử và Địa Lý (Địa lý)", "6A8", 2, ""),
    (26, "Vũ Thị Bích Hằng", "Lịch sử và Địa Lý (Địa lý)", "6A9", 2, ""),
    (26, "Vũ Thị Bích Hằng", "Lịch sử và Địa Lý (Địa lý)", "6A10", 2, ""),
    (26, "Vũ Thị Bích Hằng", "Lịch sử và Địa Lý (Địa lý)", "6A11", 2, ""),
    (26, "Vũ Thị Bích Hằng", "Lịch sử và Địa Lý (Địa lý)", "7A7", 1, ""),
    (26, "Vũ Thị Bích Hằng", "Lịch sử và Địa Lý (Địa lý)", "7A8", 1, ""),
    (26, "Vũ Thị Bích Hằng", "Lịch sử và Địa Lý (Địa lý)", "7A9", 1, ""),
    (26, "Vũ Thị Bích Hằng", "Lịch sử và Địa Lý (Địa lý)", "7A10", 1, ""),
    (26, "Vũ Thị Bích Hằng", "Lịch sử và Địa Lý (Địa lý)", "7A11", 1, ""),
    (26, "Vũ Thị Bích Hằng", "Công nghệ", "6A7", 1, ""),
    (26, "Vũ Thị Bích Hằng", "Công nghệ", "6A8", 1, ""),
    (26, "Vũ Thị Bích Hằng", "Công nghệ", "6A9", 1, ""),
    (26, "Vũ Thị Bích Hằng", "Công nghệ", "6A10", 1, ""),
    # 27. Chu Diệu Linh (CN 9A11)
    (27, "Chu Diệu Linh", "Lịch sử và Địa Lý (Địa lý)", "9A8", 1, ""),
    (27, "Chu Diệu Linh", "Lịch sử và Địa Lý (Địa lý)", "9A9", 1, ""),
    (27, "Chu Diệu Linh", "Lịch sử và Địa Lý (Địa lý)", "9A10", 1, ""),
    (27, "Chu Diệu Linh", "Lịch sử và Địa Lý (Địa lý)", "9A11", 1, ""),
    (27, "Chu Diệu Linh", "Lịch sử và Địa Lý (Địa lý)", "9A12", 1, ""),
    (27, "Chu Diệu Linh", "Lịch sử và Địa Lý (Địa lý)", "8A7", 2, ""),
    (27, "Chu Diệu Linh", "Lịch sử và Địa Lý (Địa lý)", "8A8", 2, ""),
    (27, "Chu Diệu Linh", "Lịch sử và Địa Lý (Địa lý)", "8A9", 2, ""),
    (27, "Chu Diệu Linh", "Lịch sử và Địa Lý (Địa lý)", "8A10", 2, ""),
    (27, "Chu Diệu Linh", "Lịch sử và Địa Lý (Địa lý)", "8A11", 2, ""),
    # 28. Trương Thị Thu Hương (CN 7A11)
    (28, "Trương Thị Thu Hương", "Khoa học tự nhiên (Vật lý)", "9A8", 2, ""),
    (28, "Trương Thị Thu Hương", "Khoa học tự nhiên (Vật lý)", "9A9", 2, ""),
    (28, "Trương Thị Thu Hương", "Khoa học tự nhiên (Vật lý)", "9A10", 2, ""),
    (28, "Trương Thị Thu Hương", "Khoa học tự nhiên (Vật lý)", "9A11", 2, ""),
    (28, "Trương Thị Thu Hương", "Khoa học tự nhiên (Vật lý)", "9A12", 2, ""),
    (28, "Trương Thị Thu Hương", "Khoa học tự nhiên (Vật lý)", "7A7", 1, ""),
    (28, "Trương Thị Thu Hương", "Khoa học tự nhiên (Vật lý)", "7A8", 1, ""),
    (28, "Trương Thị Thu Hương", "Khoa học tự nhiên (Vật lý)", "7A9", 1, ""),
    (28, "Trương Thị Thu Hương", "Khoa học tự nhiên (Vật lý)", "7A10", 1, ""),
    (28, "Trương Thị Thu Hương", "Khoa học tự nhiên (Vật lý)", "7A11", 1, ""),
    # 29. Đỗ Thị Thu Hiền
    (29, "Đỗ Thị Thu Hiền", "Khoa học tự nhiên (Vật lý)", "6A7", 2, ""),
    (29, "Đỗ Thị Thu Hiền", "Khoa học tự nhiên (Vật lý)", "6A8", 2, ""),
    (29, "Đỗ Thị Thu Hiền", "Khoa học tự nhiên (Vật lý)", "6A9", 2, ""),
    (29, "Đỗ Thị Thu Hiền", "Khoa học tự nhiên (Vật lý)", "6A10", 2, ""),
    (29, "Đỗ Thị Thu Hiền", "Khoa học tự nhiên (Vật lý)", "6A11", 2, ""),
    (29, "Đỗ Thị Thu Hiền", "Khoa học tự nhiên (Vật lý)", "8A7", 1, ""),
    (29, "Đỗ Thị Thu Hiền", "Khoa học tự nhiên (Vật lý)", "8A8", 1, ""),
    (29, "Đỗ Thị Thu Hiền", "Khoa học tự nhiên (Vật lý)", "8A9", 1, ""),
    (29, "Đỗ Thị Thu Hiền", "Khoa học tự nhiên (Vật lý)", "8A10", 1, ""),
    (29, "Đỗ Thị Thu Hiền", "Khoa học tự nhiên (Vật lý)", "8A11", 1, ""),
    (29, "Đỗ Thị Thu Hiền", "Hoạt động trải nghiệm, hướng nghiệp", "6A8", 3, ""),
    (29, "Đỗ Thị Thu Hiền", "Hoạt động trải nghiệm, hướng nghiệp", "8A11", 1, "Dạy chung 1 tiết HĐTN 8A11"),
    # 30. Lê Văn Hậu
    (30, "Lê Văn Hậu", "Giáo dục thể chất", "6A7", 2, ""),
    (30, "Lê Văn Hậu", "Giáo dục thể chất", "6A8", 2, ""),
    (30, "Lê Văn Hậu", "Giáo dục thể chất", "6A9", 2, ""),
    (30, "Lê Văn Hậu", "Giáo dục thể chất", "6A10", 2, ""),
    (30, "Lê Văn Hậu", "Giáo dục thể chất", "6A11", 2, ""),
    (30, "Lê Văn Hậu", "Giáo dục thể chất", "7A7", 2, ""),
    (30, "Lê Văn Hậu", "Giáo dục thể chất", "7A8", 2, ""),
    (30, "Lê Văn Hậu", "Giáo dục thể chất", "7A9", 2, ""),
    (30, "Lê Văn Hậu", "Giáo dục thể chất", "7A10", 2, ""),
    (30, "Lê Văn Hậu", "Giáo dục thể chất", "7A11", 2, ""),
    # 31. GV môn GDTC
    (31, "GV môn GDTC", "Giáo dục thể chất", "8A7", 2, ""),
    (31, "GV môn GDTC", "Giáo dục thể chất", "8A8", 2, ""),
    (31, "GV môn GDTC", "Giáo dục thể chất", "8A9", 2, ""),
    (31, "GV môn GDTC", "Giáo dục thể chất", "8A10", 2, ""),
    (31, "GV môn GDTC", "Giáo dục thể chất", "9A8", 2, ""),
    (31, "GV môn GDTC", "Giáo dục thể chất", "9A9", 2, ""),
    (31, "GV môn GDTC", "Giáo dục thể chất", "9A10", 2, ""),
    (31, "GV môn GDTC", "Giáo dục thể chất", "9A11", 2, ""),
    (31, "GV môn GDTC", "Giáo dục thể chất", "9A12", 2, ""),
    # 32. Bùi Đức Anh (Trường chính)
    (32, "Bùi Đức Anh", "Giáo dục thể chất", "8A11", 2, "Trường chính"),
    # 33. Lý Thị Duyên (CN 6A10)
    (33, "Lý Thị Duyên", "Nghệ thuật (Mỹ thuật)", "6A7", 1, ""),
    (33, "Lý Thị Duyên", "Nghệ thuật (Mỹ thuật)", "6A8", 1, ""),
    (33, "Lý Thị Duyên", "Nghệ thuật (Mỹ thuật)", "6A9", 1, ""),
    (33, "Lý Thị Duyên", "Nghệ thuật (Mỹ thuật)", "6A10", 1, ""),
    (33, "Lý Thị Duyên", "Nghệ thuật (Mỹ thuật)", "6A11", 1, ""),
    (33, "Lý Thị Duyên", "Nghệ thuật (Mỹ thuật)", "8A7", 1, ""),
    (33, "Lý Thị Duyên", "Nghệ thuật (Mỹ thuật)", "8A8", 1, ""),
    (33, "Lý Thị Duyên", "Nghệ thuật (Mỹ thuật)", "8A9", 1, ""),
    (33, "Lý Thị Duyên", "Nghệ thuật (Mỹ thuật)", "8A10", 1, ""),
    (33, "Lý Thị Duyên", "Nghệ thuật (Mỹ thuật)", "8A11", 1, ""),
    (33, "Lý Thị Duyên", "Nghệ thuật (Mỹ thuật)", "9A8", 1, ""),
    (33, "Lý Thị Duyên", "Nghệ thuật (Mỹ thuật)", "9A9", 1, ""),
    (33, "Lý Thị Duyên", "Nghệ thuật (Mỹ thuật)", "9A10", 1, ""),
    (33, "Lý Thị Duyên", "Nghệ thuật (Mỹ thuật)", "9A11", 1, ""),
    (33, "Lý Thị Duyên", "Nghệ thuật (Mỹ thuật)", "9A12", 1, ""),
    # 34. Lê Thị Minh Anh (HĐK)
    (34, "Lê Thị Minh Anh", "Nghệ thuật (Mỹ thuật)", "7A7", 1, ""),
    (34, "Lê Thị Minh Anh", "Nghệ thuật (Mỹ thuật)", "7A8", 1, ""),
    (34, "Lê Thị Minh Anh", "Nghệ thuật (Mỹ thuật)", "7A9", 1, ""),
    (34, "Lê Thị Minh Anh", "Nghệ thuật (Mỹ thuật)", "7A10", 1, ""),
    # 35. Ngô Thuý Vân (TPT)
    (35, "Ngô Thuý Vân", "Nghệ thuật (Mỹ thuật)", "7A11", 1, "Dạy TPT"),
    # 36. Nguyễn Thị Luyến
    (36, "Nguyễn Thị Luyến", "Nghệ thuật (Âm nhạc)", "6A7", 1, ""),
    (36, "Nguyễn Thị Luyến", "Nghệ thuật (Âm nhạc)", "6A8", 1, ""),
    (36, "Nguyễn Thị Luyến", "Nghệ thuật (Âm nhạc)", "6A9", 1, ""),
    (36, "Nguyễn Thị Luyến", "Nghệ thuật (Âm nhạc)", "6A10", 1, ""),
    (36, "Nguyễn Thị Luyến", "Nghệ thuật (Âm nhạc)", "6A11", 1, ""),
    (36, "Nguyễn Thị Luyến", "Nghệ thuật (Âm nhạc)", "7A7", 1, ""),
    (36, "Nguyễn Thị Luyến", "Nghệ thuật (Âm nhạc)", "7A8", 1, ""),
    (36, "Nguyễn Thị Luyến", "Nghệ thuật (Âm nhạc)", "7A9", 1, ""),
    (36, "Nguyễn Thị Luyến", "Nghệ thuật (Âm nhạc)", "7A10", 1, ""),
    (36, "Nguyễn Thị Luyến", "Nghệ thuật (Âm nhạc)", "7A11", 1, ""),
    (36, "Nguyễn Thị Luyến", "Nghệ thuật (Âm nhạc)", "8A7", 1, ""),
    (36, "Nguyễn Thị Luyến", "Nghệ thuật (Âm nhạc)", "8A8", 1, ""),
    (36, "Nguyễn Thị Luyến", "Nghệ thuật (Âm nhạc)", "8A9", 1, ""),
    (36, "Nguyễn Thị Luyến", "Nghệ thuật (Âm nhạc)", "8A10", 1, ""),
    (36, "Nguyễn Thị Luyến", "Nghệ thuật (Âm nhạc)", "8A11", 1, ""),
    (36, "Nguyễn Thị Luyến", "Nghệ thuật (Âm nhạc)", "9A8", 1, ""),
    (36, "Nguyễn Thị Luyến", "Nghệ thuật (Âm nhạc)", "9A9", 1, ""),
    (36, "Nguyễn Thị Luyến", "Nghệ thuật (Âm nhạc)", "9A10", 1, ""),
    (36, "Nguyễn Thị Luyến", "Nghệ thuật (Âm nhạc)", "9A11", 1, ""),
    (36, "Nguyễn Thị Luyến", "Nghệ thuật (Âm nhạc)", "9A12", 1, ""),
    # 37. Nguyễn Đình Tiến (CN 9A8)
    (37, "Nguyễn Đình Tiến", "Tiếng Anh", "9A8", 3, ""),
    (37, "Nguyễn Đình Tiến", "Tiếng Anh", "9A9", 3, ""),
    (37, "Nguyễn Đình Tiến", "Tiếng Anh", "7A7", 3, ""),
    (37, "Nguyễn Đình Tiến", "Tiếng Anh", "7A8", 3, ""),
    (37, "Nguyễn Đình Tiến", "Tiếng Anh", "7A9", 3, ""),
    # 38. Phan Thị Minh (CN 6A9, TPCM)
    (38, "Phan Thị Minh", "Tiếng Anh", "6A9", 3, ""),
    (38, "Phan Thị Minh", "Tiếng Anh", "6A10", 3, ""),
    (38, "Phan Thị Minh", "Tiếng Anh", "9A10", 3, ""),
    (38, "Phan Thị Minh", "Tiếng Anh", "9A11", 3, ""),
    (38, "Phan Thị Minh", "Hoạt động trải nghiệm, hướng nghiệp", "6A9", 3, ""),
    # 39. Trần Đức Ngân (CN 8A11)
    (39, "Trần Đức Ngân", "Tiếng Anh", "8A7", 3, ""),
    (39, "Trần Đức Ngân", "Tiếng Anh", "8A9", 3, ""),
    (39, "Trần Đức Ngân", "Tiếng Anh", "8A10", 3, ""),
    (39, "Trần Đức Ngân", "Tiếng Anh", "8A11", 3, ""),
    (39, "Trần Đức Ngân", "Tiếng Anh", "9A12", 3, ""),
    (39, "Trần Đức Ngân", "Hoạt động trải nghiệm, hướng nghiệp", "8A11", 2, "Dạy 2 tiết HĐTN 8A11 (Cô Hiền dạy 1t)"),
    # 40. Nguyễn Thị Hồng Minh (HĐ235)
    (40, "Nguyễn Thị Hồng Minh", "Tiếng Anh", "8A8", 3, ""),
    (40, "Nguyễn Thị Hồng Minh", "Tiếng Anh", "6A7", 2, "Dạy 2 tiết TA 6A7 (Cô Sa dạy 1t)"),
    (40, "Nguyễn Thị Hồng Minh", "Tiếng Anh", "6A8", 2, "Dạy 2 tiết TA 6A8 (Cô Sa dạy 1t)"),
    (40, "Nguyễn Thị Hồng Minh", "Tiếng Anh", "6A11", 3, ""),
    (40, "Nguyễn Thị Hồng Minh", "Tiếng Anh", "7A10", 3, ""),
    (40, "Nguyễn Thị Hồng Minh", "Tiếng Anh", "7A11", 3, ""),
    (40, "Nguyễn Thị Hồng Minh", "Hoạt động trải nghiệm, hướng nghiệp", "7A9", 3, ""),
    # 41. Bàng Như Quỳnh (Trường chính)
    (41, "Bàng Như Quỳnh", "Công nghệ", "6A11", 1, "Trường chính"),
]


def create_school_db():
    print(f"Creating database at {DB_PATH}...")
    if os.path.exists(DB_PATH):
        try:
            os.remove(DB_PATH)
        except Exception as e:
            print(f"Warning removing old DB: {e}")

    os.makedirs(os.path.dirname(DB_PATH), exist_ok=True)
    conn = db.get_connection(DB_PATH)
    conn.execute("PRAGMA foreign_keys = OFF")
    existing_tables = [row[0] for row in conn.execute("SELECT name FROM sqlite_master WHERE type='table' AND name NOT LIKE 'sqlite_%'").fetchall()]
    for tbl in existing_tables:
        conn.execute(f"DROP TABLE IF EXISTS {tbl}")
    conn.commit()
    conn.execute("PRAGMA foreign_keys = ON")
    db.init_db(conn)

    # Set metadata
    repo.set_meta(conn, "school_name", "THCS Đỗ Cận - Phân hiệu")
    repo.set_meta(conn, "base_cap", "19")
    repo.set_meta(conn, "min_floor", "15")

    # Insert classes
    class_id_map = {}
    for name, sort_order in CLASSES_DATA:
        cid = repo.upsert_class(conn, name, sort_order)
        class_id_map[name] = cid

    # Insert subjects
    subj_id_map = {}
    for name, role_code, sort_order in SUBJECTS_DATA:
        sid = repo.upsert_subject(conn, name, role_code, sort_order)
        subj_id_map[name] = sid

    # Insert teachers
    teacher_id_map = {}
    for stt, name, specialty, gvcn_class, role, reduction, thuc_day, tong, thua_thieu in TEACHERS_RAW:
        is_gvcn = 1 if gvcn_class else 0
        tid = repo.upsert_teacher(
            conn, name, role=role, must_monday=1, is_gvcn=is_gvcn,
            reduction_override=reduction if reduction else None
        )
        teacher_id_map[name] = tid

    # Insert frame templates for all 20 classes
    for name, cid in class_id_map.items():
        repo.set_frame_template(conn, cid, morning_periods=5, afternoon_periods=3, allow_saturday=False, study_sunday=False)

    # 1. Sum up periods per (subject, class) across ASSIGNMENTS_FLAT
    # to handle co-teaching (e.g. 6A7 Tiếng Anh: Cô Sa 1t + Cô Minh 2t = 3t; 8A11 HĐTN: Cô Hiền 1t + Thầy Ngân 2t = 3t)
    periods_by_sc = {}
    for item in ASSIGNMENTS_FLAT:
        stt, teacher_name, subj_name, cls_name, periods, note = item
        if cls_name not in class_id_map or subj_name not in subj_id_map:
            continue
        cid = class_id_map[cls_name]
        sid = subj_id_map[subj_name]
        periods_by_sc[(sid, cid)] = periods_by_sc.get((sid, cid), 0) + periods

    for (sid, cid), p in periods_by_sc.items():
        repo.set_periods_per_week(conn, sid, cid, "C", p)
        repo.set_periods_per_week(conn, sid, cid, "L", p)

    # 2. Insert primary timetable assignments for the 16 standard subjects
    # For co-teaching classes:
    # 6A7, 6A8 Tiếng Anh -> Cô Nguyễn Thị Hồng Minh (3t trên TKB, gồm 1t BGH Cô Sa dạy chung)
    # 8A11 HĐTN -> Thầy Trần Đức Ngân (3t trên TKB, gồm 1t Cô Hiền dạy chung)
    for item in ASSIGNMENTS_FLAT:
        stt, teacher_name, subj_name, cls_name, periods, note = item
        if cls_name not in class_id_map or subj_name not in subj_id_map:
            continue
        cid = class_id_map[cls_name]
        sid = subj_id_map[subj_name]
        tid = teacher_id_map[teacher_name]

        if teacher_name == "Đỗ Thị Sa":
            continue
        if teacher_name == "Đỗ Thị Thu Hiền" and cls_name == "8A11" and "nghiệm" in subj_name:
            continue

        repo.set_assignment(conn, sid, cid, tid)

    # 3. Populate 35-week curriculum directly from the school's real weekly periods
    # (Setting Week 1 to 35 with the accurate distribution of THCS Đỗ Cận - Phân hiệu)
    print("Populating 35-week curriculum from accurate school periods...")
    weekly_entries = []
    for w in range(1, 36):
        for (sid, cid), p in periods_by_sc.items():
            weekly_entries.append((sid, cid, w, p))
    repo.bulk_set_weekly_curriculum(conn, weekly_entries)
    print(f"Populated weekly curriculum: {len(weekly_entries)} records for 35 weeks.")

    conn.commit()
    conn.close()
    print("Database populated successfully.")


def create_excel_workbook():
    print(f"Creating Excel workbook at {EXCEL_PATH}...")
    template_path = os.path.join("io_excel", "export_template.xlsm")
    if os.path.exists(template_path):
        wb = openpyxl.load_workbook(template_path)
    else:
        wb = openpyxl.Workbook()
        wb.remove(wb.active)

    def _clear_values(ws, first_data_row: int = 2) -> None:
        for row in ws.iter_rows(min_row=first_data_row, max_row=ws.max_row):
            for cell in row:
                cell.value = None

    class_names = [c[0] for c in CLASSES_DATA]
    n_classes = len(class_names)

    # Compute matrices with combined periods for co-teaching
    pc_matrix = {}
    periods_matrix = {}
    for item in ASSIGNMENTS_FLAT:
        stt, teacher_name, subj_name, cls_name, periods, note = item
        if teacher_name == "Đỗ Thị Sa":
            continue
        if teacher_name == "Đỗ Thị Thu Hiền" and cls_name == "8A11" and "nghiệm" in subj_name:
            continue
        pc_matrix[(subj_name, cls_name)] = teacher_name

    for item in ASSIGNMENTS_FLAT:
        stt, teacher_name, subj_name, cls_name, periods, note = item
        periods_matrix[(subj_name, cls_name)] = periods_matrix.get((subj_name, cls_name), 0) + periods

    # 1. PhanCong
    if "PhanCong" in wb.sheetnames:
        ws_pc = wb["PhanCong"]
        _clear_values(ws_pc, first_data_row=2)
        ws_pc.cell(2, 1, "Môn \\ Lớp")
        for i, cname in enumerate(class_names):
            ws_pc.cell(2, 2 + i, cname)
        code_col = 2 + n_classes + 1
        ws_pc.cell(2, code_col, "MÃ VAI TRÒ")
        for s_idx, (sname, rcode, _) in enumerate(SUBJECTS_DATA, start=3):
            ws_pc.cell(s_idx, 1, sname)
            for i, cname in enumerate(class_names):
                ws_pc.cell(s_idx, 2 + i, pc_matrix.get((sname, cname), ""))
            ws_pc.cell(s_idx, code_col, rcode)

    # 2. SoTiet
    if "SoTiet" in wb.sheetnames:
        ws_st = wb["SoTiet"]
        _clear_values(ws_st, first_data_row=2)
        ws_st.cell(2, 1, "Môn \\ Lớp")
        for i, cname in enumerate(class_names):
            ws_st.cell(2, 2 + i, f"{cname} C")
        odd_start = 2 + n_classes + 1
        for i, cname in enumerate(class_names):
            ws_st.cell(2, odd_start + i, f"{cname} L")
        for s_idx, (sname, _, _) in enumerate(SUBJECTS_DATA, start=3):
            ws_st.cell(s_idx, 1, sname)
            for i, cname in enumerate(class_names):
                p = periods_matrix.get((sname, cname), 0)
                ws_st.cell(s_idx, 2 + i, p)
                ws_st.cell(s_idx, odd_start + i, p)

    # 3. DinhMuc_GV
    if "DinhMuc_GV" in wb.sheetnames:
        ws_dm = wb["DinhMuc_GV"]
        _clear_values(ws_dm, first_data_row=3)
        ws_dm.cell(1, 11, "Chức vụ")
        ws_dm.cell(1, 12, "Giảm")
        ws_dm.cell(1, 14, "Chuẩn:")
        ws_dm.cell(1, 15, 19)
        ws_dm.cell(1, 17, "Sàn tối thiểu:")
        ws_dm.cell(1, 18, 15)
        ws_dm.cell(2, 1, "Tên GV")
        ws_dm.cell(2, 2, "Chức vụ")
        ws_dm.cell(2, 8, "Đi T2 (1/0)")
        ws_dm.cell(2, 9, "GVCN (1/0)")
        for r, t in enumerate(TEACHERS_RAW):
            row = 3 + r
            stt, name, spec, gvcn, role, reduction, thuc_day, tong, diff = t
            ws_dm.cell(row, 1, name)
            ws_dm.cell(row, 2, role)
            ws_dm.cell(row, 8, 1)
            ws_dm.cell(row, 9, 1 if gvcn else 0)

        role_reductions = [
            ("Hiệu trưởng", 17), ("Phó hiệu trưởng", 15), ("Tổ trưởng", 3),
            ("Tổ phó", 1), ("GVCN", 4), ("Tổng phụ trách", 8), ("Thư ký HĐ", 2)
        ]
        for r, (r_name, r_red) in enumerate(role_reductions):
            ws_dm.cell(2 + r, 11, r_name)
            ws_dm.cell(2 + r, 12, r_red)

    # 3.1 Clear GV_Ban placeholders
    if "GV_Ban" in wb.sheetnames:
        _clear_values(wb["GV_Ban"], first_data_row=2)

    # 4. Khung & TKB_Nhap
    if "Khung" in wb.sheetnames and "TKB_Nhap" in wb.sheetnames:
        ws_khung = wb["Khung"]
        ws_nh = wb["TKB_Nhap"]
        _clear_values(ws_khung, first_data_row=2)
        _clear_values(ws_nh, first_data_row=2)

        weekdays = [2, 3, 4, 5, 6, 7]
        weekday_names = {2: 'Thứ 2', 3: 'Thứ 3', 4: 'Thứ 4', 5: 'Thứ 5', 6: 'Thứ 6', 7: 'Thứ 7', 8: 'Chủ Nhật'}
        for i, wd in enumerate(weekdays):
            ws_nh.cell(1, 4 + i, weekday_names[wd])
            ws_khung.cell(1, 4 + i, weekday_names[wd])
        ws_nh.cell(1, 4 + len(weekdays), weekday_names[8])
        ws_khung.cell(1, 4 + len(weekdays), weekday_names[8])

        row_idx = 2
        for cname in class_names:
            sessions = [('S', 1), ('S', 2), ('S', 3), ('S', 4), ('S', 5), ('C', 1), ('C', 2), ('C', 3)]
            for sess, per in sessions:
                ws_nh.cell(row_idx, 1, cname)
                ws_nh.cell(row_idx, 2, sess)
                ws_nh.cell(row_idx, 3, per)
                ws_khung.cell(row_idx, 1, cname)
                ws_khung.cell(row_idx, 2, sess)
                ws_khung.cell(row_idx, 3, per)
                for i, wd in enumerate(weekdays):
                    col = 4 + i
                    if wd <= 6:
                        ws_khung.cell(row_idx, col, 'x')
                row_idx += 1

    # 5. TuanConfig
    if "TuanConfig" in wb.sheetnames:
        ws_tc = wb["TuanConfig"]
        ws_tc.cell(1, 1, "SEED HIEN TAI ->")
        ws_tc.cell(1, 2, 0)
        ws_tc.cell(2, 1, "TUAN HIEN TAI (C=chan/L=le) ->")
        ws_tc.cell(2, 2, "C")

    # Styles
    font_title = Font(name="Calibri", size=14, bold=True, color="1F4E78")
    font_subtitle = Font(name="Calibri", size=11, italic=True)
    font_header = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    font_bold = Font(name="Calibri", size=10, bold=True)
    font_regular = Font(name="Calibri", size=10)

    fill_navy = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    fill_sub_header = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")
    fill_light_blue = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    fill_zebra = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")

    thin_border_side = Side(style="thin", color="D9D9D9")
    cell_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)

    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")

    # ==========================================
    # SHEET: Du_Lieu_Phan_Cong (Bảng phẳng / Unpivot - 1 trường dữ liệu chuẩn)
    # ==========================================
    if "Du_Lieu_Phan_Cong" in wb.sheetnames:
        ws_flat = wb["Du_Lieu_Phan_Cong"]
        ws_flat.delete_rows(1, ws_flat.max_row)
    else:
        ws_flat = wb.create_sheet(title="Du_Lieu_Phan_Cong")
    ws_flat.views.sheetView[0].showGridLines = True

    ws_flat["A1"] = "BẢNG DỮ LIỆU PHÂN CÔNG CHUYÊN MÔN CHUẨN HÓA (DẠNG BẢNG PHẲNG / 1 DÒNG 1 PHÂN CÔNG)"
    ws_flat["A1"].font = font_title
    ws_flat["A2"] = "TRƯỜNG THCS ĐỖ CẬN - PHÂN HIỆU | NĂM HỌC 2026 - 2027 (Áp dụng từ 07/09/2026)"
    ws_flat["A2"].font = font_subtitle

    headers_flat = [
        "STT GV", "Họ và tên giáo viên", "Chuyên môn", "Chức vụ / Kiêm nhiệm",
        "Chủ nhiệm lớp", "Môn giảng dạy", "Lớp giảng dạy", "Khối", "Số tiết / tuần", "Ghi chú phân công"
    ]
    for col_idx, h in enumerate(headers_flat, start=1):
        c = ws_flat.cell(4, col_idx, h)
        c.font = font_header
        c.fill = fill_navy
        c.alignment = align_center
        c.border = cell_border

    t_lookup = {t[1]: t for t in TEACHERS_RAW}
    r_idx = 5
    for item in ASSIGNMENTS_FLAT:
        stt, teacher_name, subj_name, cls_name, periods, note = item
        t_info = t_lookup.get(teacher_name, (stt, teacher_name, "", None, "", 0, 0, 0, 0))
        spec = t_info[2]
        role = t_info[4]
        gvcn = t_info[3] or ""
        grade = cls_name[:1] if cls_name[:1].isdigit() else ""

        ws_flat.cell(r_idx, 1, stt).alignment = align_center
        ws_flat.cell(r_idx, 2, teacher_name).alignment = align_left
        ws_flat.cell(r_idx, 3, spec).alignment = align_center
        ws_flat.cell(r_idx, 4, role).alignment = align_center
        ws_flat.cell(r_idx, 5, gvcn).alignment = align_center
        ws_flat.cell(r_idx, 6, subj_name).alignment = align_left
        ws_flat.cell(r_idx, 7, cls_name).alignment = align_center
        ws_flat.cell(r_idx, 8, grade).alignment = align_center
        ws_flat.cell(r_idx, 9, periods).alignment = align_right
        ws_flat.cell(r_idx, 10, note).alignment = align_left

        for c_idx in range(1, 11):
            cell = ws_flat.cell(r_idx, c_idx)
            cell.border = cell_border
            cell.font = font_regular
            if r_idx % 2 == 0:
                cell.fill = fill_zebra

        r_idx += 1

    total_row = r_idx
    ws_flat.cell(total_row, 1, "TỔNG").font = font_bold
    ws_flat.cell(total_row, 1).alignment = align_center
    ws_flat.cell(total_row, 9, f"=SUM(I5:I{total_row - 1})").font = font_bold
    ws_flat.cell(total_row, 9).alignment = align_right
    for c_idx in range(1, 11):
        cell = ws_flat.cell(total_row, c_idx)
        cell.border = cell_border
        cell.fill = fill_light_blue

    # ==========================================
    # SHEET 5: Bang_Phan_Cong_Goc (Tái tạo bố cục gốc từ ảnh)
    # ==========================================
    if "Bang_Phan_Cong_Goc" in wb.sheetnames:
        ws_goc = wb["Bang_Phan_Cong_Goc"]
        ws_goc.delete_rows(1, ws_goc.max_row)
    else:
        ws_goc = wb.create_sheet(title="Bang_Phan_Cong_Goc")
    ws_goc.views.sheetView[0].showGridLines = True

    ws_goc["A1"] = "UBND PHƯỜNG PHỔ YÊN"
    ws_goc["A1"].font = Font(name="Calibri", size=11, bold=True)
    ws_goc["A2"] = "TRƯỜNG THCS ĐỖ CẬN - PHÂN HIỆU"
    ws_goc["A2"].font = Font(name="Calibri", size=11, bold=True, underline="single")

    ws_goc["H1"] = "PHÂN CÔNG CHUYÊN MÔN HỌC KỲ I NĂM HỌC 2026-2027"
    ws_goc["H1"].font = Font(name="Calibri", size=13, bold=True)
    ws_goc["H2"] = "Thực hiện từ 7/9/2026"
    ws_goc["H2"].font = Font(name="Calibri", size=11, italic=True)

    headers_goc = [
        "STT", "Họ và tên giáo viên", "Chuyên môn", "CN Lớp", "CN S.t",
        "K.nhiệm khác", "K.nhiệm S.t",
        "Dạy môn 1", "S.t 1", "Dạy môn 2", "S.t 2", "Dạy môn 3", "S.t 3", "Dạy môn 4", "S.t 4", "Dạy môn 5", "S.t 5",
        "BDHSG / CLB", "S.t BD", "Thực dạy", "Tổng", "Thừa (+)/Thiếu (-)"
    ]
    for c_idx, h in enumerate(headers_goc, start=1):
        c = ws_goc.cell(4, c_idx, h)
        c.font = font_header
        c.fill = fill_navy
        c.alignment = align_center
        c.border = cell_border

    # Group assignments by teacher
    teacher_assignments_map = {}
    for item in ASSIGNMENTS_FLAT:
        stt, tname, subj, cls, p, note = item
        if tname not in teacher_assignments_map:
            teacher_assignments_map[tname] = []
        teacher_assignments_map[tname].append((subj, cls, p, note))

    for r_idx, t in enumerate(TEACHERS_RAW, start=5):
        stt, name, spec, gvcn, role, reduction, thuc_day, tong, diff = t
        ws_goc.cell(r_idx, 1, stt).alignment = align_center
        ws_goc.cell(r_idx, 2, name).alignment = align_left
        ws_goc.cell(r_idx, 3, spec).alignment = align_center
        ws_goc.cell(r_idx, 4, gvcn or "").alignment = align_center
        ws_goc.cell(r_idx, 5, 4 if gvcn else "").alignment = align_center
        ws_goc.cell(r_idx, 6, role if role != "GVCN" else "").alignment = align_center
        ws_goc.cell(r_idx, 7, reduction if role != "GVCN" and reduction else "").alignment = align_center

        # Group teaching items
        items = teacher_assignments_map.get(name, [])
        reg_items = [it for it in items if not it[0].startswith("BDHSG")]
        bd_items = [it for it in items if it[0].startswith("BDHSG")]

        # Group regular by subject if multi-class
        subj_grouped = {}
        for subj, cls, p, note in reg_items:
            if subj not in subj_grouped:
                subj_grouped[subj] = []
            subj_grouped[subj].append((cls, p))

        slot_idx = 0
        for subj, cls_list in subj_grouped.items():
            if slot_idx >= 5:
                break
            # Format text: e.g. "Toán 6A7,8"
            cls_str = ",".join(c[0] for c in cls_list)
            tot_p = sum(c[1] for c in cls_list)
            txt = f"{subj} {cls_str}"

            col_m = 8 + slot_idx * 2
            col_p = col_m + 1
            ws_goc.cell(r_idx, col_m, txt).alignment = align_left
            ws_goc.cell(r_idx, col_p, tot_p).alignment = align_center
            slot_idx += 1

        # BDHSG
        if bd_items:
            ws_goc.cell(r_idx, 18, bd_items[0][0]).alignment = align_left
            ws_goc.cell(r_idx, 19, bd_items[0][2]).alignment = align_center

        ws_goc.cell(r_idx, 20, thuc_day).alignment = align_right
        ws_goc.cell(r_idx, 21, tong).alignment = align_right
        ws_goc.cell(r_idx, 22, diff if diff is not None else "").alignment = align_right

        for c_idx in range(1, 23):
            cell = ws_goc.cell(r_idx, c_idx)
            cell.border = cell_border
            cell.font = font_regular
            if r_idx % 2 == 1:
                cell.fill = fill_zebra

    # Auto-fit all sheets
    for sheet in wb.worksheets:
        for col in sheet.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val = cell.value
                if val:
                    lines = str(val).split("\n")
                    max_len = max(max_len, max(len(l) for l in lines))
            sheet.column_dimensions[col_letter].width = min(max(max_len + 3, 10), 38)

    wb.save(EXCEL_PATH)
    print(f"Workbook saved successfully to {EXCEL_PATH}.")


if __name__ == "__main__":
    create_school_db()
    create_excel_workbook()
    print("ALL DONE.")
