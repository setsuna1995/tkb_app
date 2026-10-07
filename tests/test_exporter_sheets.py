import io

import openpyxl
from core.models import ROLE_THUONG, ClassRoom, Subject, Teacher
from io_excel.exporter import export_class_sheets_xlsx, export_teacher_sheets_xlsx

CLASSES = [ClassRoom(2, "6/2", 2), ClassRoom(1, "6A1", 1)]
SUBJECTS = [Subject(1, "Toán", ROLE_THUONG), Subject(2, "Văn", ROLE_THUONG)]
TEACHERS = [Teacher(10, "Nguyễn Thị Hà"), Teacher(11, "Nguyễn Thị Hà"), Teacher(12, "GV Không Dạy")]
ASSIGN = {(1, 1): 10, (2, 1): 11, (1, 2): 10}
CELLS = {(1, 2, "S", 1): 1, (1, 2, "S", 2): 2, (2, 3, "C", 2): 1, (2, 2, "S", 1): None}


def _load(data):
    return openpyxl.load_workbook(io.BytesIO(data))


def test_class_workbook_one_sheet_per_class_in_order():
    wb = _load(export_class_sheets_xlsx(CELLS, CLASSES, SUBJECTS, TEACHERS, ASSIGN, title_suffix=" — Tuần 6"))
    assert wb.sheetnames == ["6A1", "6-2"]            # sort_order, "/" sanitised
    ws = wb["6A1"]
    assert ws["A1"].value == "Thời khóa biểu lớp 6A1 — Tuần 6"
    assert [c.value for c in ws[2]][:3] == ["Buổi", "Tiết", "Thứ 2"]
    assert ws["C3"].value == "Toán\nNguyễn Thị Hà"     # Thứ 2, Sáng, tiết 1
    assert ws.max_row == 2 + 5                          # 6A1 has no afternoon -> morning rows only
    assert wb["6-2"].max_row == 2 + 10                  # 6/2 has an afternoon period


def test_teacher_workbook_dedupes_names_and_skips_idle():
    wb = _load(export_teacher_sheets_xlsx(CELLS, CLASSES, SUBJECTS, TEACHERS, ASSIGN))
    assert sorted(wb.sheetnames) == ["Nguyễn Thị Hà", "Nguyễn Thị Hà (2)"]
    texts = {c.value for ws in wb for row in ws.iter_rows() for c in row if c.value}
    assert {"6A1 (Toán)", "6/2 (Toán)", "6A1 (Văn)"} <= texts


def test_sheet_name_truncated_to_31():
    wb = _load(export_class_sheets_xlsx({(1, 2, "S", 1): 1}, [ClassRoom(1, "L" * 40, 1)], SUBJECTS, TEACHERS, {(1, 1): 10}))
    assert len(wb.sheetnames[0]) == 31


def test_empty_input_still_valid_workbook():
    assert _load(export_teacher_sheets_xlsx({}, [], SUBJECTS, [], {})).sheetnames == ["Trống"]
